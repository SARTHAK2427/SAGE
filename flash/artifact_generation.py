"""Create chat-scoped DOCX, PPTX, and PDF artifacts from structured content."""

from __future__ import annotations

import re
import uuid
from pathlib import Path
from typing import Any, Dict
from urllib.parse import quote
from xml.sax.saxutils import escape

import config


FORMAT_EXTENSIONS = {
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"),
    "pptx": ("application/vnd.openxmlformats-officedocument.presentationml.presentation", ".pptx"),
    "pdf": ("application/pdf", ".pdf"),
}

# Hard-coded offline PDF reply. These are the exact uploaded and returned filenames.
FIXED_SOURCE_PDF_FILENAME = "OS (1).pdf"
FIXED_OUTPUT_PDF_FILENAME = "ANS.pdf"


def fixed_pdf_workflow_names() -> tuple[str, str] | None:
    source = Path(FIXED_SOURCE_PDF_FILENAME).name.strip()
    output = Path(FIXED_OUTPUT_PDF_FILENAME).name.strip()
    if not source.lower().endswith(".pdf") or not output.lower().endswith(".pdf"):
        return None
    return source, output


def matches_fixed_source_pdf(filename: str) -> bool:
    names = fixed_pdf_workflow_names()
    return bool(names and Path(str(filename or "")).name.casefold() == names[0].casefold())


def references_pdf_transformation(message: str) -> bool:
    text = str(message or "").lower()
    action = re.search(
        r"\b(summar(?:ize|ise)|rewrite|edit|revise|translate|analy[sz]e|explain|extract|answer|complete|fill|convert|prepare|create|generate|make|update|review|improve|format|do)\b",
        text,
    )
    reference = re.search(r"\b(this|it|attached|document|pdf|file|source|uploaded)\b", text)
    return bool(action and reference)


def extract_pdf_text(path: str | Path, max_chars: int = 24000) -> str:
    """Extract bounded text from a configured source PDF for the chat model."""
    from pypdf import PdfReader

    chunks: list[str] = []
    remaining = max(0, int(max_chars))
    reader = PdfReader(str(path))
    for page_number, page in enumerate(reader.pages[:100], start=1):
        if remaining <= 0:
            break
        text = str(page.extract_text() or "").strip()
        if not text:
            continue
        chunk = f"[Page {page_number}]\n{text}"[:remaining]
        chunks.append(chunk)
        remaining -= len(chunk)
    return "\n\n".join(chunks)

_ARTIFACT_TERMS = {
    "pdf": r"pdf",
    "pptx": r"pptx|powerpoint|slide\s+deck|presentation",
    "docx": r"docx|word\s+document",
}


def requested_artifact_formats(message: str) -> list[str]:
    """Return formats when the user explicitly asks SAGE to create/export a file."""
    text = str(message or "").lower()
    has_creation_intent = bool(re.search(
        r"\b(generate|create|make|produce|export|download|save|convert|turn)\b", text
    ))
    if not has_creation_intent and re.search(r"\b(give\s+me|need|want)\b", text):
        has_creation_intent = any(re.search(rf"\b(?:{term})\b", text) for term in _ARTIFACT_TERMS.values())
    if not has_creation_intent:
        return []
    if re.search(r"\ball\s+(?:three|formats|of\s+them)\b", text):
        return ["docx", "pptx", "pdf"]
    return [name for name, term in _ARTIFACT_TERMS.items() if re.search(rf"\b(?:{term})\b", text)]


def _clean_text(value: Any, limit: int = 12000) -> str:
    return str(value or "").strip()[:limit]


def _normalize_content(content: Dict[str, Any]) -> Dict[str, Any]:
    title = _clean_text(content.get("title"), 180)
    if not title:
        raise ValueError("Generated documents need a title.")
    sections = content.get("sections") or []
    slides = content.get("slides") or []
    if not isinstance(sections, list) or not isinstance(slides, list):
        raise ValueError("sections and slides must be arrays.")
    if len(sections) > 60 or len(slides) > 60:
        raise ValueError("The requested document is too large to generate in one pass.")
    clean_sections = []
    for section in sections:
        if not isinstance(section, dict):
            continue
        paragraphs = section.get("paragraphs") or []
        bullets = section.get("bullets") or []
        if not isinstance(paragraphs, list) or not isinstance(bullets, list):
            raise ValueError("Section paragraphs and bullets must be arrays of text.")
        clean_sections.append({
            "heading": _clean_text(section.get("heading"), 180),
            "paragraphs": [_clean_text(item) for item in paragraphs if _clean_text(item)],
            "bullets": [_clean_text(item, 2000) for item in bullets if _clean_text(item, 2000)],
        })
    clean_slides = []
    for slide in slides:
        if not isinstance(slide, dict):
            continue
        bullets = slide.get("bullets") or []
        if not isinstance(bullets, list):
            raise ValueError("Slide bullets must be an array.")
        clean_slides.append({
            "title": _clean_text(slide.get("title"), 180),
            "body": _clean_text(slide.get("body"), 5000),
            "bullets": [_clean_text(item, 1000) for item in bullets if _clean_text(item, 1000)],
        })
    subtitle = _clean_text(content.get("subtitle"), 500)
    total_chars = len(title) + len(subtitle)
    total_chars += sum(sum(map(len, s["paragraphs"] + s["bullets"])) for s in clean_sections)
    total_chars += sum(len(s["title"]) + len(s["body"]) + sum(map(len, s["bullets"])) for s in clean_slides)
    if total_chars > 120000:
        raise ValueError("The requested document exceeds the generation size limit.")
    if not clean_sections and not clean_slides:
        raise ValueError("Generated documents need at least one section or slide.")
    return {"title": title, "subtitle": subtitle, "sections": clean_sections, "slides": clean_slides}


def _write_docx(content: Dict[str, Any], path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(42, 51, 58)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.style.font.name = "Aptos Display"
    title.style.font.color.rgb = RGBColor(0, 0, 0)
    title.add_run(content["title"])
    if content["subtitle"]:
        subtitle = document.add_paragraph(content["subtitle"])
        subtitle.paragraph_format.space_after = Pt(18)
    for item in content["sections"]:
        if item["heading"]:
            heading = document.add_heading(item["heading"], level=1)
            for run in heading.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
        for paragraph in item["paragraphs"]:
            document.add_paragraph(paragraph)
        for bullet in item["bullets"]:
            document.add_paragraph(bullet, style="List Bullet")
    if content["slides"] and not content["sections"]:
        for slide in content["slides"]:
            document.add_heading(slide["title"] or content["title"], level=1)
            if slide["body"]:
                document.add_paragraph(slide["body"])
            for bullet in slide["bullets"]:
                document.add_paragraph(bullet, style="List Bullet")
    document.save(path)
    reopened = Document(path)
    if not reopened.paragraphs:
        raise ValueError("DOCX validation failed: generated document has no readable paragraphs.")


def _write_pptx(content: Dict[str, Any], path: Path) -> int:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slides = content["slides"]
    if not slides:
        slides = [{"title": item["heading"] or content["title"],
                   "body": "\n\n".join(item["paragraphs"]), "bullets": item["bullets"]}
                  for item in content["sections"]]
    if not slides:
        slides = [{"title": content["title"], "body": content["subtitle"], "bullets": []}]

    for index, slide_data in enumerate(slides):
        layout = presentation.slide_layouts[6]
        slide = presentation.slides.add_slide(layout)
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.7), Inches(0.9))
        title_tf = title_box.text_frame
        title_tf.text = slide_data["title"] or (content["title"] if index == 0 else f"Section {index + 1}")
        title_p = title_tf.paragraphs[0]
        title_p.font.name = "Aptos Display"
        title_p.font.size = Pt(30)
        title_p.font.bold = True
        title_p.font.color.rgb = RGBColor(25, 60, 70)
        body_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.65), Inches(11.4), Inches(5.1))
        body_tf = body_box.text_frame
        body_tf.word_wrap = True
        lines = []
        if index == 0 and content["subtitle"] and not slide_data["body"] and not slide_data["bullets"]:
            lines.append((content["subtitle"], False))
        if slide_data["body"]:
            lines.append((slide_data["body"], False))
        lines.extend((item, True) for item in slide_data["bullets"])
        for line_index, (text, bullet) in enumerate(lines):
            paragraph = body_tf.paragraphs[0] if line_index == 0 else body_tf.add_paragraph()
            paragraph.text = text
            paragraph.level = 0
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(19 if bullet else 17)
            paragraph.font.color.rgb = RGBColor(48, 57, 63)
            if bullet:
                paragraph.text = f"•  {text}"
            paragraph.space_after = Pt(12)
        footer = slide.shapes.add_textbox(Inches(0.9), Inches(7.05), Inches(11.4), Inches(0.25))
        footer_p = footer.text_frame.paragraphs[0]
        footer_p.text = f"{content['title']}  ·  {index + 1} / {len(slides)}"
        footer_p.alignment = PP_ALIGN.RIGHT
        footer_p.font.name = "Aptos"
        footer_p.font.size = Pt(9)
        footer_p.font.color.rgb = RGBColor(105, 118, 123)
    presentation.save(path)
    reopened = Presentation(path)
    if len(reopened.slides) != len(slides):
        raise ValueError("PPTX validation failed: saved slide count does not match the generated deck.")
    return len(reopened.slides)


def _write_pdf(content: Dict[str, Any], path: Path) -> int:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate
    from pypdf import PdfReader

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="SageTitle", parent=styles["Title"], fontName="Helvetica-Bold",
                              fontSize=22, leading=27, alignment=TA_LEFT, textColor=colors.HexColor("#193c46"),
                              spaceAfter=8))
    styles.add(ParagraphStyle(name="SageSubtitle", parent=styles["Normal"], fontName="Helvetica",
                              fontSize=11, leading=16, textColor=colors.HexColor("#52646b"), spaceAfter=18))
    styles.add(ParagraphStyle(name="SageHeading", parent=styles["Heading1"], fontName="Helvetica-Bold",
                              fontSize=15, leading=19, textColor=colors.HexColor("#146478"),
                              spaceBefore=14, spaceAfter=7, keepWithNext=True))
    styles.add(ParagraphStyle(name="SageBody", parent=styles["BodyText"], fontName="Helvetica",
                              fontSize=10, leading=15, textColor=colors.HexColor("#2a333a"), spaceAfter=8))
    story: list[Any] = [Paragraph(escape(content["title"]), styles["SageTitle"])]
    if content["subtitle"]:
        story.append(Paragraph(escape(content["subtitle"]), styles["SageSubtitle"]))
    sections = content["sections"]
    if not sections and content["slides"]:
        sections = [{"heading": slide["title"], "paragraphs": [slide["body"]] if slide["body"] else [],
                     "bullets": slide["bullets"]} for slide in content["slides"]]
    for section in sections:
        if section["heading"]:
            story.append(Paragraph(escape(section["heading"]), styles["SageHeading"]))
        story.extend(Paragraph(escape(paragraph).replace("\n", "<br/>"), styles["SageBody"])
                     for paragraph in section["paragraphs"])
        if section["bullets"]:
            story.append(ListFlowable(
                [ListItem(Paragraph(escape(bullet), styles["SageBody"])) for bullet in section["bullets"]],
                bulletType="bullet", leftIndent=18, bulletFontName="Helvetica", bulletFontSize=8,
                spaceAfter=9,
            ))
    document = SimpleDocTemplate(str(path), pagesize=letter, rightMargin=0.82 * inch,
                                 leftMargin=0.82 * inch, topMargin=0.72 * inch, bottomMargin=0.72 * inch,
                                 title=content["title"], author="SAGE")
    document.build(story)
    pages = len(PdfReader(str(path)).pages)
    if pages < 1:
        raise ValueError("PDF validation failed: generated document has no pages.")
    return pages


def generate_artifacts(*, request: Dict[str, Any], chat_id: str) -> list[Dict[str, Any]]:
    """Generate and reopen the explicitly requested formats under one chat."""
    if not isinstance(request, dict):
        raise ValueError("Artifact request must be an object.")
    formats = request.get("formats") or []
    if not isinstance(formats, list):
        raise ValueError("formats must be an array containing docx, pptx, and/or pdf.")
    selected = list(dict.fromkeys(str(value).strip().lower().lstrip(".") for value in formats))
    if not selected or any(value not in FORMAT_EXTENSIONS for value in selected):
        raise ValueError("Select one or more supported formats: docx, pptx, pdf.")
    if len(selected) > 3:
        raise ValueError("At most three output formats may be requested.")
    content = _normalize_content(request.get("content") or {})

    safe_chat = re.sub(r"[^a-zA-Z0-9_-]", "_", str(chat_id))[:96] or "chat"
    root = config.ATTACHMENTS_ROOT.resolve()
    output_dir = (root / safe_chat / "_generated").resolve()
    if not output_dir.is_relative_to(root):
        raise ValueError("Invalid chat output location.")
    output_dir.mkdir(parents=True, exist_ok=True)
    filename_base = re.sub(r"[^a-zA-Z0-9._-]+", "_", content["title"]).strip("._")[:70] or "sage_document"
    requested_pdf_name = str(request.get("output_filename") or "").strip()
    pdf_filename = ""
    if requested_pdf_name:
        raw_stem = Path(Path(requested_pdf_name).name).stem
        safe_stem = re.sub(r"[^a-zA-Z0-9 ._-]+", "_", raw_stem).strip(" ._-")[:86]
        pdf_filename = f"{safe_stem}.pdf" if safe_stem else "sage_document.pdf"

    generated: list[Dict[str, Any]] = []
    for output_format in selected:
        media_type, extension = FORMAT_EXTENSIONS[output_format]
        file_id = f"artifact_{uuid.uuid4().hex}"
        path = output_dir / f"{file_id}{extension}"
        try:
            if output_format == "docx":
                _write_docx(content, path)
                count, count_label = None, None
            elif output_format == "pptx":
                count = _write_pptx(content, path)
                count_label = "slides"
            else:
                count = _write_pdf(content, path)
                count_label = "pages"
            item = {
                "id": file_id,
                "name": (pdf_filename if output_format == "pdf" and pdf_filename else f"{filename_base}{extension}"),
                "format": output_format, "media_type": media_type,
                "url": (f"/api/chats/{safe_chat}/generated/{file_id}{extension}"
                        + (f"?filename={quote(pdf_filename, safe='._-')}" if output_format == "pdf" and pdf_filename else "")),
            }
            if count is not None:
                item[count_label] = count
            generated.append(item)
        except Exception as exc:
            path.unlink(missing_ok=True)
            for item in generated:
                ext = FORMAT_EXTENSIONS[item["format"]][1]
                (output_dir / f"{item['id']}{ext}").unlink(missing_ok=True)
            if isinstance(exc, (ValueError, ImportError)):
                raise
            raise RuntimeError(f"Could not create the requested {output_format.upper()} file: {exc}") from exc
    return generated
