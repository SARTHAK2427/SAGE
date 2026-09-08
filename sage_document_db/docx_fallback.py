"""
sage_document_db/docx_fallback.py

Lightweight fallback parser for DOCX files using python-docx.
Invoked ONLY if Docling fails on a (sanitized) DOCX file.

Truthful Reporting Contract:
- parser_name is explicitly set to "docx_fallback" (NOT docling).
- parser_version is set to docx.__version__.
- Explicit warning is appended: does not claim full-fidelity parsing;
  documents that only paragraphs and tables were recovered.
"""

from __future__ import annotations
from pathlib import Path
from typing import Optional

try:
    import docx
    DOCX_VERSION = getattr(docx, "__version__", "unknown")
except ImportError:
    docx = None
    DOCX_VERSION = None

from .models import NormalizedDocument, NormalizedElement


def parse_docx_fallback(
    working_path: Path | str,
    original_source_path: Path | str,
    doc_id: str,
    original_sha: str,
    original_size: int,
    docling_error: Optional[str] = None,
) -> NormalizedDocument:
    """Parse a DOCX file using python-docx when Docling fails.

    Preserves text paragraphs and tables. Truthfully marks the document
    as parsed via docx_fallback.
    """
    if docx is None:
        raise ImportError(
            "python-docx is not installed; cannot perform fallback DOCX parsing."
        )

    working_path = Path(working_path)
    original_source_path = Path(original_source_path)

    doc_obj = docx.Document(str(working_path))

    elements: list[NormalizedElement] = []
    txt_count = 0
    table_count = 0
    global_order = 0

    # 1. Extract paragraphs
    for p in doc_obj.paragraphs:
        text = p.text
        if text is None:
            continue
        # We record paragraph even if empty or whitespace to preserve sequence
        txt_count += 1
        el_id = f"txt_{txt_count:06d}"
        elements.append(
            NormalizedElement(
                id=el_id,
                type="text",
                subtype="paragraph",
                order=global_order,
                text=text,
            )
        )
        global_order += 1

    # 2. Extract tables
    for tbl in doc_obj.tables:
        rows_data: list[list[str]] = []
        for row in tbl.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            rows_data.append(row_cells)

        if not rows_data:
            continue

        table_count += 1
        el_id = f"table_{table_count:06d}"

        # Simple markdown table representation
        md_lines = []
        if rows_data:
            headers = rows_data[0]
            md_lines.append("| " + " | ".join(headers) + " |")
            md_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
            for r in rows_data[1:]:
                md_lines.append("| " + " | ".join(r) + " |")
        table_md = "\n".join(md_lines)

        elements.append(
            NormalizedElement(
                id=el_id,
                type="table",
                subtype="table",
                order=global_order,
                text=table_md,
                table_data={
                    "rows": rows_data,
                    "markdown": table_md,
                },
            )
        )
        global_order += 1

    # 3. Set reading order chain
    for i, el in enumerate(elements):
        el.previous_element = elements[i - 1].id if i > 0 else None
        el.next_element = elements[i + 1].id if i + 1 < len(elements) else None

    # 4. Construct warnings with truthful fallback notice
    warnings = [
        f"Docling parsing failed ({docling_error or 'conversion error'}); parsed using python-docx fallback. "
        "Recovered text paragraphs and tables; visual formatting, embedded images, and layout coordinates are omitted."
    ]

    return NormalizedDocument(
        doc_id=doc_id,
        source_path=str(original_source_path),
        original_name=original_source_path.name,
        extension=".docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        sha256=original_sha,
        file_size_bytes=original_size,
        parser_name="docx_fallback",
        parser_version=DOCX_VERSION,
        elements=elements,
        warnings=warnings,
    )
