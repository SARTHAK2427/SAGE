"""Smoke checks for chat-scoped office and PDF artifact generation."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from flash.artifact_generation import (
    extract_pdf_text, generate_artifacts, matches_fixed_source_pdf,
    references_pdf_transformation, requested_artifact_formats,
)


class ArtifactGenerationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.output_root = Path(self.temp.name)
        self.patch_root = patch("flash.artifact_generation.config.ATTACHMENTS_ROOT", self.output_root)
        self.patch_root.start()
        self.addCleanup(self.patch_root.stop)
        self.content = {
            "title": "Project Overview",
            "subtitle": "A concise generated artifact.",
            "sections": [{
                "heading": "Purpose",
                "paragraphs": ["This paragraph verifies that generated text survives serialization."],
                "bullets": ["First deliverable", "Second deliverable"],
            }],
            "slides": [{
                "title": "Project Overview",
                "body": "A concise generated artifact.",
                "bullets": ["First deliverable", "Second deliverable"],
            }],
        }

    def test_generates_and_reopens_all_three_requested_formats(self):
        files = generate_artifacts(
            request={"formats": ["docx", "pptx", "pdf"], "content": self.content},
            chat_id="test_chat",
        )
        self.assertEqual({item["format"] for item in files}, {"docx", "pptx", "pdf"})
        for item in files:
            self.assertTrue(item["url"].startswith("/api/chats/test_chat/generated/"))
            self.assertTrue((self.output_root / "test_chat" / "_generated" / f"{item['id']}.{item['format']}").is_file())

        docx_file = next(item for item in files if item["format"] == "docx")
        doc = Document(self.output_root / "test_chat" / "_generated" / f"{docx_file['id']}.docx")
        self.assertIn("This paragraph verifies", " ".join(p.text for p in doc.paragraphs))

        pptx_file = next(item for item in files if item["format"] == "pptx")
        deck = Presentation(self.output_root / "test_chat" / "_generated" / f"{pptx_file['id']}.pptx")
        self.assertEqual(len(deck.slides), 1)

        pdf_file = next(item for item in files if item["format"] == "pdf")
        pdf = PdfReader(self.output_root / "test_chat" / "_generated" / f"{pdf_file['id']}.pdf")
        self.assertEqual(len(pdf.pages), pdf_file["pages"])
        self.assertIn("Project Overview", pdf.pages[0].extract_text())

    def test_creates_only_the_requested_format(self):
        files = generate_artifacts(
            request={"formats": ["pdf"], "content": self.content},
            chat_id="pdf_only",
        )
        self.assertEqual([item["format"] for item in files], ["pdf"])

    def test_pdf_can_use_a_fixed_download_filename(self):
        files = generate_artifacts(
            request={"formats": ["pdf"], "output_filename": "../Fixed Report.pdf", "content": self.content},
            chat_id="fixed_name",
        )
        self.assertEqual(files[0]["name"], "Fixed Report.pdf")
        self.assertIn("?filename=Fixed%20Report.pdf", files[0]["url"])
        pdf_path = self.output_root / "fixed_name" / "_generated" / f"{files[0]['id']}.pdf"
        self.assertIn("This paragraph verifies", extract_pdf_text(pdf_path))

    def test_source_pdf_filename_is_exact_and_configurable(self):
        from unittest.mock import patch

        with patch("flash.artifact_generation.FIXED_SOURCE_PDF_FILENAME", "Source File.pdf"), \
             patch("flash.artifact_generation.FIXED_OUTPUT_PDF_FILENAME", "Finished File.pdf"):
            self.assertTrue(matches_fixed_source_pdf("Source File.pdf"))
            self.assertTrue(matches_fixed_source_pdf("source file.PDF"))
            self.assertFalse(matches_fixed_source_pdf("different.pdf"))
        self.assertTrue(references_pdf_transformation("summarize this"))
        self.assertFalse(references_pdf_transformation("hello there"))

    def test_rejects_unsupported_format_without_creating_files(self):
        with self.assertRaises(ValueError):
            generate_artifacts(
                request={"formats": ["html"], "content": self.content},
                chat_id="invalid",
            )
        self.assertFalse((self.output_root / "invalid").exists())

    def test_detects_explicit_requested_formats(self):
        self.assertEqual(requested_artifact_formats("Generate a PDF explaining this."), ["pdf"])
        self.assertEqual(requested_artifact_formats("Make a PPTX presentation about the results."), ["pptx"])
        self.assertEqual(requested_artifact_formats("Generate all three formats."), ["docx", "pptx", "pdf"])
        self.assertEqual(requested_artifact_formats("Can you explain how PDFs are created?"), [])
        self.assertFalse(matches_fixed_source_pdf("YOUR_INPUT_PDF_NAME.pdf"))


if __name__ == "__main__":
    unittest.main()
